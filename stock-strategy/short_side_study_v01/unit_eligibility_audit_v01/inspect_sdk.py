"""Read-only static SDK/document inventory. Never imports or executes the SDK."""
import argparse
import hashlib
import json
import re
import struct
import subprocess
from pathlib import Path


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def fixed_versions(path):
    """Read VS_FIXEDFILEINFO signature/version words; not assembly execution."""
    data = path.read_bytes()
    offset = data.find(b'\xbd\x04\xef\xfe')
    if offset < 0 or offset + 52 > len(data):
        return None
    values = struct.unpack_from('<13I', data, offset)
    if values[1] != 0x10000:
        return None
    version = lambda a, b: '.'.join(map(str, (a >> 16, a & 65535, b >> 16, b & 65535)))
    return {'fixed_file_info_offset': offset, 'file_version': version(values[2], values[3]),
            'product_version': version(values[4], values[5])}


def inspect(directory):
    from openpyxl import load_workbook
    from pypdf import PdfReader
    result = {'directory': str(directory), 'execution': 'STATIC_ONLY_NO_SDK_IMPORT',
              'files': [], 'workbook_rows': [], 'pdf_hits': [], 'doc_excerpts': []}
    names = ('YuantaSparkAPI.dll', 'NX_StockInfo_Model.dll', 'NX_Trade_TW_Model.dll',
             'YuantaCAPIDLL.dll', 'YuantaCAPIDLL64.dll', 'YuantaSparkAPI.deps.json',
             'YuantaSparkAPI.runtimeconfig.json', 'YSendOrder.py')
    for name in names:
        p = directory / name
        if p.is_file():
            result['files'].append({'path': str(p), 'sha256': digest(p), 'bytes': p.stat().st_size,
                                    'fixed_versions': fixed_versions(p) if p.suffix == '.dll' else None})
    deps = json.loads((directory / 'YuantaSparkAPI.deps.json').read_text(encoding='utf-8-sig'))
    result['declared_packages'] = {k: v for k, v in deps['libraries'].items()
                                   if k.startswith(('YuantaSparkAPI/', 'NX_StockInfo_Model/', 'NX_Trade_TW_Model/'))}
    workbook = directory / 'IO_Doc/FunctionList.xlsx'
    result['files'].append({'path': str(workbook), 'sha256': digest(workbook)})
    wb = load_workbook(workbook, read_only=True, data_only=False)
    result['workbook_sheets'] = []
    for sheet in wb:
        result['workbook_sheets'].append({'name': sheet.title, 'rows': sheet.max_row, 'columns': sheet.max_column})
        for cells in sheet.iter_rows():
            text = ' '.join(str(c.value) for c in cells if c.value is not None)
            if re.search('FiveTick|Vol[1-5]|MarketType|IndexFlag|OrderQty|成交量|委託量|零股|TWSE|TWOTC|單位', text, re.I):
                result['workbook_rows'].append({'sheet': sheet.title,
                    'range': f'A{cells[0].row}:{cells[-1].column_letter}{cells[0].row}',
                    'cells': {c.coordinate: str(c.value) for c in cells if c.value is not None}})
    wb.close()
    for pdf in sorted(directory.glob('*.pdf')):
        reader = PdfReader(pdf)
        result['files'].append({'path': str(pdf), 'sha256': digest(pdf), 'pages': len(reader.pages)})
        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ''
            if re.search('IndexFlag|enumMarketType|BuyVol|SellVol|Vol1|OrderQty|單位|五檔', text, re.I):
                result['pdf_hits'].append({'path': str(pdf), 'page': i + 1, 'text': text})
    for name in ('SubscribeFiveTickA', 'SendStockOrder', 'GetRealReport', 'GetRealReportMerge',
                 'RR_RealReport', 'RR_RealReportMerge', 'SubscribeStocktick', 'GetStockInformation'):
        path = directory / f'IO_Doc/{name}_IO_Spec.doc'
        text = subprocess.run(['textutil', '-convert', 'txt', '-stdout', str(path)],
                              check=True, capture_output=True, text=True).stdout
        result['doc_excerpts'].append({'path': str(path), 'sha256': digest(path), 'text': text})
    return result


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--sdk', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    result = inspect(args.sdk)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({'files': len(result['files']), 'workbook_rows': len(result['workbook_rows']),
                      'pdf_hits': len(result['pdf_hits'])}))
